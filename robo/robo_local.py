"""
robo/robo_local.py
==================
Robô monitor de pasta local para arquivos TOTALE (CSV/Excel).

Detecção automática de BASE pelo nome do arquivo:
    • "ABCDM" no nome → Base NET-ABCDM
    • "SPO"   no nome → Base NET-LESTE
    • "GRS"   no nome → Base NET-GUARULHOS

Leitura Excel à prova de:
    • OptionError: No such keys(s): 'io.excel.zip.reader'
    • Arquivo bloqueado (OneDrive/SharePoint ainda sincronizando)
    • openpyxl ausente / engine incompatível
    • .xlsx corrompido (zip inválido / download parcial)

API pública: renderizar_robo_local(etl_fn, gsheets_fn, ...)
"""

from __future__ import annotations

import os
import tempfile
import zipfile
from collections.abc import Callable
from datetime import datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st

try:
    from pandas.errors import OptionError as PandasOptionError
except ImportError:  # pandas < 1.0
    PandasOptionError = KeyError  # type: ignore[misc, assignment]

# ── Constantes ──────────────────────────────────────────────────────
PADROES_TOTALE: tuple[str, ...] = (
    "Atividades-*.csv",
    "Atividades-*.xlsx",
    "Atividades-*.xls",
)
EXTENSOES_VALIDAS: tuple[str, ...] = (".csv", ".xlsx", ".xls")
EXTENSOES_TEMP: tuple[str, ...] = (".tmp", ".crdownload")

MAPA_BASES_ARQUIVO: dict[str, str] = {
    "ABCDM": "NET-ABCDM",
    "SPO": "NET-LESTE",
    "GRS": "NET-GUARULHOS",
}

_VALORES_BASE_VAZIOS: frozenset[str] = frozenset(
    {"", "nan", "none", "na", "n/a", "não informado", "nao informado", "-"}
)
_TAMANHO_MINIMO_XLSX_VALIDO: int = 2_000


# ── Detecção de Base ────────────────────────────────────────────────
def detectar_base_arquivo(caminho_ou_nome: str | None) -> str | None:
    if not caminho_ou_nome:
        return None
    nome = Path(str(caminho_ou_nome)).name.upper()
    for marcador in sorted(MAPA_BASES_ARQUIVO, key=len, reverse=True):
        if marcador in nome:
            return MAPA_BASES_ARQUIVO[marcador]
    return None


def detectar_marcador_arquivo(caminho_ou_nome: str | None) -> str | None:
    if not caminho_ou_nome:
        return None
    nome = Path(str(caminho_ou_nome)).name.upper()
    for marcador in sorted(MAPA_BASES_ARQUIVO, key=len, reverse=True):
        if marcador in nome:
            return marcador
    return None


def _aplicar_base_no_dataframe(df: pd.DataFrame, base: str) -> pd.DataFrame:
    df = df.copy()
    if "BASE" not in df.columns:
        df["BASE"] = base
        return df
    serie = df["BASE"]
    vazios = serie.isna() | serie.astype(str).str.strip().str.lower().isin(
        _VALORES_BASE_VAZIOS
    )
    df.loc[vazios, "BASE"] = base
    return df


# ── Utilitários ─────────────────────────────────────────────────────
def obter_pasta_robo_padrao() -> str:
    base = Path(__file__).resolve().parent
    for cand in (base / "dados", Path.home() / "robo", Path.home() / "Downloads"):
        if cand.exists() and cand.is_dir():
            return str(cand.resolve())
    return str(Path.home() / "Downloads")


def formatar_tamanho(bytes_: int) -> str:
    if bytes_ < 1024:
        return f"{bytes_} B"
    if bytes_ < 1024 * 1024:
        return f"{bytes_ / 1024:.1f} KB"
    return f"{bytes_ / (1024 * 1024):.2f} MB"


def obter_metadados_arquivo(caminho: str | None) -> dict[str, str]:
    if not caminho or not os.path.exists(caminho):
        return {
            "nome": Path(caminho).name if caminho else "Desconhecido",
            "ext": Path(caminho).suffix.lower() if caminho else "",
            "tamanho": "—",
            "modificado": "—",
            "caminho_completo": caminho or "",
        }
    try:
        p = Path(caminho)
        stat = p.stat()
        return {
            "nome": p.name,
            "ext": p.suffix.lower(),
            "tamanho": formatar_tamanho(stat.st_size),
            "modificado": datetime.fromtimestamp(stat.st_mtime).strftime(
                "%d/%m/%Y %H:%M:%S"
            ),
            "caminho_completo": str(p.resolve()),
        }
    except Exception:
        return {
            "nome": Path(caminho).name if caminho else "Desconhecido",
            "ext": Path(caminho).suffix.lower() if caminho else "",
            "tamanho": "—",
            "modificado": "—",
            "caminho_completo": caminho or "",
        }


def _stat_basico(caminho: str) -> tuple[float | None, int | None]:
    try:
        st_ = os.stat(caminho)
        return float(st_.st_mtime), int(st_.st_size)
    except Exception:
        return None, None


def _toast_dedup(mensagem: str, icon: str = "✅") -> None:
    assinatura = f"{icon}:{mensagem}"
    if st.session_state.get("_robo_ultimo_toast") != assinatura:
        st.session_state["_robo_ultimo_toast"] = assinatura
        st.toast(mensagem, icon=icon)


def _arquivo_esta_bloqueado(caminho: str) -> bool:
    try:
        with open(caminho, "rb") as fh:
            fh.read(1)
        return False
    except (PermissionError, OSError):
        return True
    except Exception:
        return False


def _normalizar_df_excel(
    df: pd.DataFrame | dict[str, pd.DataFrame] | None,
) -> pd.DataFrame:
    if df is None:
        return pd.DataFrame()
    if isinstance(df, dict):
        return next(iter(df.values())) if df else pd.DataFrame()
    return df


# ── Leitura Excel robusta (corrige io.excel.zip.reader) ─────────────
class ArquivoBloqueadoError(Exception):
    """Arquivo em gravação/sincronização — aguardar o próximo ciclo."""


def _ler_excel_via_openpyxl_nativo(
    raw: bytes, sheet_name: str | int
) -> pd.DataFrame:
    """
    Bypass total do registry de engines do pandas (não usa io.excel.*).
    Robusto contra: worksheet None, linhas de tamanho irregular e
    cabeçalhos duplicados/vazios.
    """
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        # 1) Seleção segura da worksheet (wb.active pode ser None)
        ws = None
        if isinstance(sheet_name, int):
            if 0 <= sheet_name < len(wb.worksheets):
                ws = wb.worksheets[sheet_name]
        elif sheet_name in wb.sheetnames:
            ws = wb[str(sheet_name)]

        if ws is None:
            ws = wb.active
        if ws is None:
            ws = wb.worksheets[0] if wb.worksheets else None
        if ws is None:
            return pd.DataFrame()

        rows = ws.iter_rows(values_only=True)

        header_raw = next(rows, None)
        if header_raw is None:
            return pd.DataFrame()

        # 2) Normaliza cabeçalho: sem None e sem nomes duplicados
        header: list[str] = []
        vistos: dict[str, int] = {}
        for i, c in enumerate(header_raw):
            nome = str(c).strip() if c is not None else ""
            if not nome:
                nome = f"col_{i}"
            if nome in vistos:
                vistos[nome] += 1
                nome = f"{nome}.{vistos[nome]}"
            else:
                vistos[nome] = 0
            header.append(nome)

        n_cols = len(header)

        # 3) Ajusta cada linha ao tamanho do header (corta ou completa)
        data: list[list[str]] = []
        for row in rows:
            valores = ["" if c is None else str(c) for c in row]
            if len(valores) < n_cols:
                valores.extend([""] * (n_cols - len(valores)))
            elif len(valores) > n_cols:
                valores = valores[:n_cols]
            data.append(valores)

        return pd.DataFrame(data, columns=header)
    finally:
        wb.close()


def _ler_excel_robusto(
    caminho: str,
    sheet_name: str | int | None = 0,
) -> pd.DataFrame:
    """
    Lê .xlsx/.xls copiando os bytes primeiro (libera lock do OneDrive).

    Estratégia em cascata:
      1. Temp file com sufixo .xlsx + engine='openpyxl'
         → evita OptionError io.excel.zip.reader
      2. BytesIO com .name='arquivo.xlsx' + engine='openpyxl'
      3. engine='calamine' (python-calamine)
      4. openpyxl.load_workbook nativo (sem pandas)
    """
    planilha: str | int = 0 if sheet_name is None else sheet_name

    try:
        with open(caminho, "rb") as fh:
            raw = fh.read()
    except PermissionError as e:
        raise ArquivoBloqueadoError(
            f"'{Path(caminho).name}' está bloqueado (OneDrive/Excel)."
        ) from e

    if len(raw) < _TAMANHO_MINIMO_XLSX_VALIDO:
        raise ValueError(
            f"Arquivo muito pequeno ({len(raw)} bytes) — download incompleto."
        )

    # .xlsx é zip: valida integridade antes de gastar engine
    if Path(caminho).suffix.lower() == ".xlsx":
        bio_chk = BytesIO(raw)
        if not zipfile.is_zipfile(bio_chk):
            raise ValueError(
                "Conteúdo não é um .xlsx válido (zip corrompido ou extensão errada)."
            )
        bio_chk.seek(0)
        with zipfile.ZipFile(bio_chk) as zf:
            ruim = zf.testzip()
            if ruim is not None:
                raise ValueError(f".xlsx corrompido no membro interno: {ruim}")

    erros: list[str] = []

    def _tentar(rotulo: str, fn: Callable[[], pd.DataFrame]) -> pd.DataFrame | None:
        try:
            df = _normalizar_df_excel(fn())
            if df is None or df.empty:
                erros.append(f"{rotulo}: planilha vazia")
                return None
            return df
        except PandasOptionError as e:
            erros.append(f"{rotulo}: OptionError {e}")
            return None
        except ImportError as e:
            erros.append(f"{rotulo}: engine ausente ({e})")
            return None
        except ArquivoBloqueadoError:
            raise
        except PermissionError as e:
            raise ArquivoBloqueadoError(str(e)) from e
        except Exception as e:
            erros.append(f"{rotulo}: {type(e).__name__}: {e}")
            return None

    # 1) Temp .xlsx — força a extensão correta no pandas (FIX PRINCIPAL)
    def _via_temp() -> pd.DataFrame:
        fd, tmp = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        try:
            with open(tmp, "wb") as out:
                out.write(raw)
            return pd.read_excel(tmp, sheet_name=planilha, dtype=str, engine="openpyxl")
        finally:
            try:
                os.unlink(tmp)
            except OSError:
                pass

    df = _tentar("openpyxl/temp.xlsx", _via_temp)
    if df is not None:
        return df

    # 2) BytesIO com .name — pandas deixa de inferir ext='zip'
    def _via_bytesio() -> pd.DataFrame:
        bio = BytesIO(raw)
        bio.name = "arquivo.xlsx"
        return pd.read_excel(bio, sheet_name=planilha, dtype=str, engine="openpyxl")

    df = _tentar("openpyxl/BytesIO", _via_bytesio)
    if df is not None:
        return df

    # 3) calamine (não usa io.excel.zip.reader)
    def _via_calamine() -> pd.DataFrame:
        bio = BytesIO(raw)
        bio.name = "arquivo.xlsx"
        return pd.read_excel(bio, sheet_name=planilha, dtype=str, engine="calamine")

    df = _tentar("calamine", _via_calamine)
    if df is not None:
        return df

    # 4) openpyxl puro — zero pandas engine registry
    df = _tentar(
        "openpyxl nativo",
        lambda: _ler_excel_via_openpyxl_nativo(raw, planilha),
    )
    if df is not None:
        return df

    dica = ""
    joined = " | ".join(erros)
    if "openpyxl" in joined.lower() and (
        "import" in joined.lower() or "engine ausente" in joined.lower()
    ):
        dica = " → instale com: pip install openpyxl"
    raise RuntimeError(
        f"{joined}{dica}" if erros else "Falha desconhecida ao ler Excel."
    )


# ── Busca ───────────────────────────────────────────────────────────
def _listar_candidatos_com_stat(
    pasta: str,
    padroes: tuple[str, ...] = PADROES_TOTALE,
) -> list[tuple[str, float, int]]:
    p = Path(pasta)
    if not p.exists() or not p.is_dir():
        return []

    encontrados: set[Path] = set()
    for padrao in padroes:
        encontrados.update(p.glob(padrao))

    if not encontrados:
        for ext in EXTENSOES_VALIDAS:
            encontrados.update(p.glob(f"*Atividades*{ext}"))

    if not encontrados:
        for marcador in MAPA_BASES_ARQUIVO:
            for ext in EXTENSOES_VALIDAS:
                encontrados.update(p.glob(f"*{marcador}*{ext}"))

    def _eh_valido(arq: Path) -> bool:
        nome = arq.name.lower()
        if nome.endswith(EXTENSOES_TEMP) or nome.startswith("~$"):
            return False
        return arq.is_file()

    candidatos: list[tuple[str, float, int]] = []
    for arq in encontrados:
        if not _eh_valido(arq):
            continue
        mtime, size = _stat_basico(str(arq))
        if mtime is None or size is None:
            continue
        candidatos.append((str(arq), mtime, size))

    candidatos.sort(key=lambda t: t[1], reverse=True)
    return candidatos


def buscar_arquivo_mais_recente(
    pasta: str,
    padroes: tuple[str, ...] = PADROES_TOTALE,
) -> tuple[str | None, float | None, int | None]:
    candidatos = _listar_candidatos_com_stat(pasta, padroes)
    if not candidatos:
        return None, None, None
    caminho, mtime, size = candidatos[0]
    return caminho, mtime, size


def mapear_arquivos_por_base(
    pasta: str,
) -> dict[str, tuple[str, float, int]]:
    por_base: dict[str, tuple[str, float, int]] = {}
    for caminho, mtime, size in _listar_candidatos_com_stat(pasta):
        base = detectar_base_arquivo(caminho)
        if base and base not in por_base:
            por_base[base] = (caminho, mtime, size)
    return por_base


# ── Leitura pública ─────────────────────────────────────────────────
def ler_arquivo_totale(
    caminho_arquivo: str,
    colunas_esperadas: list[str] | None = None,
    sheet_name: str | int | None = 0,
) -> pd.DataFrame | None:
    if not caminho_arquivo or not os.path.exists(caminho_arquivo):
        _toast_dedup(
            f"Arquivo não encontrado: {Path(caminho_arquivo).name if caminho_arquivo else '—'}",
            icon="❌",
        )
        return None

    nome_arq = Path(caminho_arquivo).name
    ext = Path(caminho_arquivo).suffix.lower()

    if _arquivo_esta_bloqueado(caminho_arquivo):
        raise ArquivoBloqueadoError(
            f"'{nome_arq}' está bloqueado (gravação/sincronização em andamento)."
        )

    if ext in (".xlsx", ".xls"):
        try:
            df = _ler_excel_robusto(caminho_arquivo, sheet_name=sheet_name)
        except ArquivoBloqueadoError:
            raise
        except Exception as e:
            st.session_state["robo_erro"] = f"{nome_arq}: {type(e).__name__} — {e}"
            _toast_dedup(f"Falha Excel: {nome_arq}", icon="⚠️")
            return None

        if df.empty:
            st.session_state["robo_erro"] = (
                f"{nome_arq}: planilha lida, porém sem linhas de dados."
            )
            _toast_dedup(f"Excel vazio: {nome_arq}", icon="⚠️")
            return None

        if colunas_esperadas and not all(c in df.columns for c in colunas_esperadas):
            faltando = [c for c in colunas_esperadas if c not in df.columns]
            st.session_state["robo_erro"] = (
                f"{nome_arq}: colunas ausentes → {', '.join(faltando)}"
            )
            _toast_dedup("Excel OK, mas colunas esperadas ausentes.", icon="⚠️")
            return None

        _toast_dedup(f"Arquivo carregado: {nome_arq}", icon="✅")
        return df

    # ── CSV ──────────────────────────────────────────────────────────
    ultimo_erro: Exception | None = None
    for enc in ("utf-8-sig", "cp1252", "latin1", "iso-8859-1", "utf-8"):
        for sep in (";", ","):
            try:
                df = pd.read_csv(
                    caminho_arquivo,
                    sep=sep,
                    encoding=enc,
                    dtype=str,
                    low_memory=False,
                    on_bad_lines="skip",
                )
                if df.shape[1] <= 3:
                    continue
                if colunas_esperadas and not all(
                    c in df.columns for c in colunas_esperadas
                ):
                    continue
                _toast_dedup(f"Arquivo carregado: {nome_arq}", icon="✅")
                return df
            except PermissionError:
                raise ArquivoBloqueadoError(
                    f"'{nome_arq}' está aberto em outro programa."
                ) from None
            except Exception as e:
                ultimo_erro = e
                continue

    detalhe = f" ({type(ultimo_erro).__name__}: {ultimo_erro})" if ultimo_erro else ""
    st.session_state["robo_erro"] = f"{nome_arq}: falha ao decodificar CSV{detalhe}"
    _toast_dedup(f"Falha ao ler: {nome_arq}", icon="⚠️")
    return None


# ── Loop (fragment) ─────────────────────────────────────────────────
@st.fragment(run_every="5s")
def _executar_verificacao_robo(
    pasta_monitorada: str,
    etl_fn: Callable[[pd.DataFrame, pd.DataFrame], pd.DataFrame] | None = None,
    gsheets_fn: Callable[[], pd.DataFrame] | None = None,
    colunas_esperadas: list[str] | None = None,
    sheet_name: str | int | None = 0,
    ciclos_estabilidade: int = 2,
) -> None:
    if not st.session_state.get("robo_ativo", True):
        return
    if st.session_state.get("robo_processando", False):
        return

    candidatos = _listar_candidatos_com_stat(pasta_monitorada)
    if not candidatos:
        return

    assinatura = tuple(candidatos)
    if st.session_state.get("robo_candidato_sig") == assinatura:
        stable = int(st.session_state.get("robo_candidato_stable", 0) or 0) + 1
    else:
        stable = 0
    st.session_state["robo_candidato_sig"] = assinatura
    st.session_state["robo_candidato_stable"] = stable

    if stable < ciclos_estabilidade:
        return
    if st.session_state.get("robo_processado_sig") == assinatura:
        return

    st.session_state["robo_processando"] = True
    try:
        dfs: list[pd.DataFrame] = []
        bases_carregadas: list[str] = []
        processados: dict[str, float] = {}
        houve_bloqueio = False

        for caminho, mtime, _size in candidatos:
            try:
                df_parcial = ler_arquivo_totale(caminho, colunas_esperadas, sheet_name)
            except ArquivoBloqueadoError as e:
                houve_bloqueio = True
                st.session_state["robo_candidato_stable"] = 0
                st.session_state["_robo_aguardando_lock"] = str(e)
                continue

            if df_parcial is None:
                continue

            base = detectar_base_arquivo(caminho)
            if base:
                df_parcial = _aplicar_base_no_dataframe(df_parcial, base)
                if base not in bases_carregadas:
                    bases_carregadas.append(base)

            dfs.append(df_parcial)
            processados[caminho] = mtime

        if houve_bloqueio and not dfs:
            return
        if not dfs:
            return

        if tuple(_listar_candidatos_com_stat(pasta_monitorada)) != assinatura:
            st.session_state["robo_candidato_stable"] = 0
            _toast_dedup("Arquivos mudaram durante a leitura. Aguardando...", icon="⚠️")
            return

        df_local = pd.concat(dfs, ignore_index=True)

        if callable(gsheets_fn):
            df_gs = gsheets_fn()
        elif isinstance(gsheets_fn, pd.DataFrame):
            df_gs = gsheets_fn
        else:
            df_gs = pd.DataFrame()

        caminho_recente, mtime_recente, _ = candidatos[0]
        st.session_state["robo_ultimo_processado_path"] = caminho_recente
        st.session_state["robo_ultimo_processado_mtime"] = mtime_recente
        st.session_state["robo_processado_sig"] = assinatura
        st.session_state["robo_arquivos_processados"] = processados
        st.session_state["robo_bases_carregadas"] = bases_carregadas
        st.session_state["robo_base_detectada"] = detectar_base_arquivo(caminho_recente)
        st.session_state["robo_hora_sucesso"] = datetime.now()
        st.session_state.pop("_robo_aguardando_lock", None)

        sufixo_bases = f" · {', '.join(bases_carregadas)}" if bases_carregadas else ""
        st.session_state["origem_dados"] = (
            f"Robô Local ({len(dfs)} arquivo(s){sufixo_bases})"
        )
        st.session_state.pop("robo_erro", None)

        if callable(etl_fn):
            resultado = etl_fn(df_local, df_gs)
            if (
                resultado is not None
                and isinstance(resultado, pd.DataFrame)
                and st.session_state.get("df_memoria") is None
            ):
                st.session_state["df_memoria"] = resultado
        else:
            st.session_state["df_memoria"] = df_local

        if bases_carregadas:
            _toast_dedup(
                f"Bases: {', '.join(bases_carregadas)} — pipeline concluído!",
                icon="🚀",
            )
        else:
            _toast_dedup("Pipeline concluído!", icon="🚀")
        st.rerun()
    except Exception as e:
        st.session_state["robo_erro"] = f"{type(e).__name__}: {e}"
        _toast_dedup(f"Erro no pipeline: {e}", icon="❌")
    finally:
        st.session_state["robo_processando"] = False


# ── UI Sidebar ──────────────────────────────────────────────────────
def renderizar_robo_local(
    etl_fn: Callable[[pd.DataFrame, pd.DataFrame], pd.DataFrame] | None = None,
    gsheets_fn: Callable[[], pd.DataFrame] | None = None,
    pasta_padrao: str | None = None,
    colunas_esperadas: list[str] | None = None,
    sheet_name: str | int | None = 0,
    ciclos_estabilidade: int = 2,
    mostrar_toggle: bool = True,
    mostrar_config: bool = True,
) -> None:
    if "robo_ativo" not in st.session_state:
        st.session_state["robo_ativo"] = True

    if "robo_pasta_alvo" not in st.session_state or not st.session_state.get(
        "robo_pasta_alvo"
    ):
        st.session_state["robo_pasta_alvo"] = pasta_padrao or obter_pasta_robo_padrao()

    if pasta_padrao and not st.session_state.get("_robo_pasta_user_set"):
        if st.session_state.get("robo_pasta_alvo") in (
            None,
            "",
            obter_pasta_robo_padrao(),
        ):
            st.session_state["robo_pasta_alvo"] = pasta_padrao

    pasta_alvo: str = str(st.session_state["robo_pasta_alvo"])

    st.sidebar.markdown(
        """
        <div style='font-size:11px;font-weight:700;color:#64748B;text-transform:uppercase;
        letter-spacing:0.8px;margin:14px 0 6px;padding-left:4px;'>
             Robô de Sincronismo
        </div>
        """,
        unsafe_allow_html=True,
    )

    if mostrar_toggle:
        ativo = st.sidebar.toggle(
            "⚡ Auto-Sincronizar",
            value=bool(st.session_state["robo_ativo"]),
            key="robo_toggle_ui",
            help=(
                "Monitora a pasta em busca de Atividades-*.csv/xlsx. "
                "Detecta a base pelo nome: SPO→Leste, GRS→Guarulhos, ABCDM→ABCDM."
            ),
        )
        st.session_state["robo_ativo"] = ativo
    else:
        ativo = bool(st.session_state["robo_ativo"])

    if mostrar_config:
        with st.sidebar.expander("⚙️ Configurar Pasta", expanded=False):
            pasta_input = st.text_input(
                "Caminho do Disco:",
                value=pasta_alvo,
                key="input_pasta_robo",
            )
            if pasta_input != pasta_alvo:
                st.session_state["_robo_pasta_user_set"] = True
            st.session_state["robo_pasta_alvo"] = pasta_input
            pasta_alvo = pasta_input

            if os.path.isdir(pasta_alvo):
                st.success("✅ Pasta válida")
                try:
                    n = len(
                        [
                            f
                            for f in os.listdir(pasta_alvo)
                            if f.lower().endswith(EXTENSOES_VALIDAS)
                        ]
                    )
                    st.caption(f"📄 {n} arquivo(s) compatível(is)")
                except Exception:
                    st.caption("⚠️ Não foi possível listar")
            else:
                st.error("❌ Pasta não existe")

    candidatos_ui = _listar_candidatos_com_stat(pasta_alvo)
    por_base = mapear_arquivos_por_base(pasta_alvo)

    if candidatos_ui:
        hora = st.session_state.get("robo_hora_sucesso")
        hs = hora.strftime("%H:%M:%S") if isinstance(hora, datetime) else "—"
        sig_ok = st.session_state.get("robo_processado_sig") == tuple(candidatos_ui)
        stable = int(st.session_state.get("robo_candidato_stable", 0) or 0)
        aguardando_lock = st.session_state.get("_robo_aguardando_lock")

        if sig_ok:
            st.sidebar.success(
                f"✅ Sincronizado às {hs}\n{len(candidatos_ui)} arquivo(s) processado(s)"
            )
        elif aguardando_lock:
            st.sidebar.info(f"⏳ {aguardando_lock}\nAguardando liberação do arquivo...")
        else:
            st.sidebar.warning(
                f"📥 {len(candidatos_ui)} arquivo(s) detectado(s)\n"
                f"Estável {stable}/{ciclos_estabilidade}"
            )

        st.sidebar.markdown(
            """
            <div style='font-size:10px;font-weight:700;color:#64748B;text-transform:uppercase;
            letter-spacing:0.6px;margin:10px 0 4px;padding-left:4px;'>
                Bases Monitoradas
            </div>
            """,
            unsafe_allow_html=True,
        )
        processados_map: dict[str, float] = (
            st.session_state.get("robo_arquivos_processados", {}) or {}
        )
        for marcador, base in MAPA_BASES_ARQUIVO.items():
            info = por_base.get(base)
            if info is None:
                st.sidebar.caption(f"⬜ {base} — aguardando arquivo com '{marcador}'")
                continue
            caminho_b, mtime_b, _ = info
            meta = obter_metadados_arquivo(caminho_b)
            ok = processados_map.get(caminho_b) == mtime_b
            icone = "✅" if ok else "📥"
            st.sidebar.caption(
                f"{icone} {base} — {meta['nome'][:30]} ({meta['tamanho']})"
            )

        sem_base = [c for c, _, _ in candidatos_ui if detectar_base_arquivo(c) is None]
        if sem_base:
            st.sidebar.caption(
                f"⚠️ {len(sem_base)} arquivo(s) sem marcador de base "
                "(SPO/GRS/ABCDM) no nome."
            )
    else:
        st.sidebar.info(
            f"{'🟢 Monitorando' if ativo else '⏸ Pausado'}: `{Path(pasta_alvo).name}`"
        )

    erro = st.session_state.get("robo_erro")
    if erro:
        st.sidebar.error(f"❌ {erro}")
        c_a, c_b = st.sidebar.columns(2)
        with c_a:
            if st.button("🧹 Limpar", key="limpar_erro_robo", use_container_width=True):
                st.session_state.pop("robo_erro", None)
                st.rerun()
        with c_b:
            if st.button(
                "🔄 Retentar", key="retry_erro_robo", use_container_width=True
            ):
                st.session_state.pop("robo_erro", None)
                st.session_state["robo_candidato_sig"] = None
                st.session_state["robo_processado_sig"] = None
                st.session_state["robo_candidato_stable"] = 0
                st.rerun()

    _executar_verificacao_robo(
        pasta_monitorada=pasta_alvo,
        etl_fn=etl_fn,
        gsheets_fn=gsheets_fn,
        colunas_esperadas=colunas_esperadas,
        sheet_name=sheet_name,
        ciclos_estabilidade=ciclos_estabilidade,
    )


# Aliases
renderizar_sidebar_robo = renderizar_robo_local
_obter_pasta_robo_padrao = obter_pasta_robo_padrao
_obter_metadados_arquivo = obter_metadados_arquivo

__all__ = [
    "EXTENSOES_VALIDAS",
    "MAPA_BASES_ARQUIVO",
    "ArquivoBloqueadoError",
    "buscar_arquivo_mais_recente",
    "detectar_base_arquivo",
    "detectar_marcador_arquivo",
    "ler_arquivo_totale",
    "mapear_arquivos_por_base",
    "obter_metadados_arquivo",
    "obter_pasta_robo_padrao",
    "renderizar_robo_local",
    "renderizar_sidebar_robo",
]
