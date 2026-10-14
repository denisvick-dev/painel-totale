"""
robo/robo_engine.py
===================
Motor unificado de IO e processamento para arquivos TOTALE (CSV/Excel).

Não importa Streamlit. Concentra toda a descoberta de arquivos, verificação de
estabilidade/assinatura, detecção de base (ABCDM, SPO, GRS), leitura resiliente
de CSV/Excel e padronização de DataFrames.
"""

from __future__ import annotations

import fnmatch
import logging
import math
import os
import tempfile
import time
import zipfile
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Literal, cast

import pandas as pd

try:
    from pandas.errors import OptionError as PandasOptionError
except ImportError:  # pandas < 1.0
    PandasOptionError = KeyError  # type: ignore[misc, assignment]

logger = logging.getLogger(__name__)

# ── Constantes ──────────────────────────────────────────────────────
EngineExcel = Literal["openpyxl", "calamine", "xlrd"]
ModoLeitura = Literal["ultimo", "ultimo_por_base", "todos"]
Candidato = tuple[str, float, int]

PADROES_TOTALE: tuple[str, ...] = (
    "Atividades-*.csv",
    "Atividades-*.xlsx",
    "Atividades-*.xls",
)
EXTENSOES_VALIDAS: tuple[str, ...] = (".csv", ".xlsx", ".xls")
EXTENSOES_TEMP: tuple[str, ...] = (
    ".tmp",
    ".crdownload",
    ".partial",
    ".download",
    ".part",
)

MAPA_BASES_ARQUIVO: dict[str, str] = {
    "ABCDM": "NET-ABCDM",
    "SPO": "NET-LESTE",
    "GRS": "NET-GUARULHOS",
}

_VALORES_AUSENTES: frozenset[str] = frozenset(
    {
        "",
        "nan",
        "none",
        "na",
        "n/a",
        "<na>",
        "nat",
        "null",
        "não informado",
        "nao informado",
        "-",
    }
)
_TAMANHO_MINIMO_XLSX: int = 2_000
_TAMANHO_MAXIMO_PADRAO: int = 80 * 1024 * 1024


class ArquivoBloqueadoError(Exception):
    """Arquivo em gravação ou sincronização. O próximo ciclo tenta de novo."""


class _BytesNomeados(BytesIO):
    """BytesIO com .name, para o pandas não inferir a extensão como zip."""

    def __init__(self, raw: bytes, nome: str) -> None:
        super().__init__(raw)
        self.name = nome


@dataclass(frozen=True)
class LeituraArquivo:
    df: pd.DataFrame | None
    erro: str | None
    motor: str | None


@dataclass(frozen=True)
class ResultadoCiclo:
    df: pd.DataFrame | None
    bases: tuple[str, ...]
    arquivos: tuple[str, ...]
    erros: tuple[str, ...]
    bloqueio: str | None
    motores: tuple[str, ...]
    duplicatas_removidas: int


# ── Texto e colunas ─────────────────────────────────────────────────
def _celula_para_texto(valor: object) -> str:
    if valor is None or valor is pd.NA:
        return ""
    if isinstance(valor, bool):
        return "True" if valor else "False"
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if math.isnan(valor):
            return ""
        if valor.is_integer():
            return str(int(valor))
        return str(valor).strip()
    texto = valor.strip() if isinstance(valor, str) else str(valor).strip()
    if texto.lower() in {"nat", "<na>", "nan", "none"}:
        return ""
    corpo = texto[:-2] if texto.endswith(".0") else ""
    if corpo and corpo.replace("-", "", 1).isdigit():
        return corpo
    return texto


def _nomes_colunas_unicos(nomes: Iterable[object]) -> list[str]:
    vistos: dict[str, int] = {}
    saida: list[str] = []
    for i, bruto in enumerate(nomes):
        nome = (
            ""
            if bruto is None
            else str(bruto).replace("\ufeff", "").replace("\xa0", " ").strip()
        )
        if not nome or nome.lower().startswith("unnamed"):
            nome = f"col_{i}"
        if nome in vistos:
            vistos[nome] += 1
            nome = f"{nome}.{vistos[nome]}"
        else:
            vistos[nome] = 0
        saida.append(nome)
    return saida


def _padronizar_coluna_texto(serie: pd.Series) -> pd.Series:
    """Converte uma coluna para texto limpo, de forma **vetorial**.

    Equivalente a aplicar :func:`_celula_para_texto` célula a célula, porém sem
    callback Python por valor. Em bases reais (120 mil linhas × 18 colunas) a
    padronização caiu de ~0,82 s para ~0,13 s (**6× mais rápida**) com saída
    idêntica — ver ``tests/test_performance_carga.py``.
    """
    if pd.api.types.is_bool_dtype(serie):
        return serie.map({True: "True", False: "False"}).astype("string")

    if pd.api.types.is_datetime64_any_dtype(serie):
        # str(Timestamp) → "AAAA-MM-DD HH:MM:SS"; mantém o formato anterior.
        return serie.dt.strftime("%Y-%m-%d %H:%M:%S").astype("string")

    if pd.api.types.is_numeric_dtype(serie):
        numerica = pd.to_numeric(serie, errors="coerce")
        texto = numerica.astype("string")
        inteiros = numerica.notna() & (numerica == numerica.round())
        if inteiros.any():
            # 12.0 → "12" (o float integral não vira texto com sufixo ".0").
            texto = texto.where(~inteiros, numerica.round().astype("Int64").astype("string"))
        return texto

    texto = serie.astype("string").str.strip()
    com_sufixo = texto.str.endswith(".0", na=False).fillna(False)
    if com_sufixo.any():
        corpo = texto.str[:-2]
        digitos = corpo.str.fullmatch(r"-?\d+", na=False).fillna(False)
        texto = texto.mask(com_sufixo & digitos, corpo)
    return texto


def _padronizar_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty and len(df.columns) == 0:
        return df.copy()
    saida = df.copy()
    saida.columns = pd.Index(_nomes_colunas_unicos(list(saida.columns)))
    for col in saida.columns:
        texto = _padronizar_coluna_texto(saida[col]).str.strip()
        saida[col] = texto.mask(texto.str.lower().isin(_VALORES_AUSENTES))
    return saida.dropna(how="all")


def _cabecalho_fraco(colunas: Sequence[object]) -> bool:
    nomes = [str(c).strip() for c in colunas]
    if not nomes:
        return True
    ruins = sum(
        1
        for n in nomes
        if not n or n.lower().startswith("unnamed") or n.lower().startswith("col_")
    )
    return ruins / len(nomes) >= 0.5


def _alinhar_colunas_esperadas(
    df: pd.DataFrame,
    esperadas: Sequence[str] | None,
) -> pd.DataFrame:
    if not esperadas:
        return df
    por_chave = {str(c).casefold(): str(c) for c in df.columns}
    renomear: dict[str, str] = {}
    for esperada in esperadas:
        if esperada in df.columns:
            continue
        atual = por_chave.get(esperada.casefold())
        if atual and atual not in renomear:
            renomear[atual] = esperada
    if not renomear:
        return df
    return df.rename(columns=renomear)


def _chaves_colunas(df: pd.DataFrame) -> set[str]:
    return {str(c).casefold() for c in df.columns}


def _contem_colunas(df: pd.DataFrame, esperadas: Sequence[str] | None) -> bool:
    if not esperadas:
        return True
    chaves = _chaves_colunas(df)
    return all(esperada.casefold() in chaves for esperada in esperadas)


def _faltam_colunas(df: pd.DataFrame, esperadas: Sequence[str] | None) -> list[str]:
    if not esperadas:
        return []
    chaves = _chaves_colunas(df)
    return [c for c in esperadas if c.casefold() not in chaves]


# ── Detecção de base ────────────────────────────────────────────────
def detectar_marcador_arquivo(caminho_ou_nome: str | Path | None) -> str | None:
    if not caminho_ou_nome:
        return None
    nome = Path(str(caminho_ou_nome)).name.upper()
    for marcador in sorted(MAPA_BASES_ARQUIVO, key=len, reverse=True):
        if marcador in nome:
            return marcador
    return None


def detectar_base_arquivo(caminho_ou_nome: str | Path | None) -> str | None:
    marcador = detectar_marcador_arquivo(caminho_ou_nome)
    if marcador is None:
        return None
    return MAPA_BASES_ARQUIVO.get(marcador)


def _canonizar_coluna_base(df: pd.DataFrame) -> pd.DataFrame:
    variantes = [
        col for col in df.columns if str(col).casefold() == "base" and col != "BASE"
    ]
    if not variantes:
        return df
    saida = df.copy()
    if "BASE" not in saida.columns:
        saida = saida.rename(columns={variantes[0]: "BASE"})
        variantes = variantes[1:]
    for col in variantes:
        texto = saida["BASE"].map(_celula_para_texto).str.strip().str.lower()
        vazios = texto.isin(_VALORES_AUSENTES)
        saida.loc[vazios, "BASE"] = saida.loc[vazios, col]
        saida = saida.drop(columns=[col])
    return saida


def _aplicar_base_no_dataframe(df: pd.DataFrame, base: str) -> pd.DataFrame:
    saida = _canonizar_coluna_base(df)
    if "BASE" not in saida.columns:
        saida["BASE"] = base
        return saida
    texto = saida["BASE"].map(_celula_para_texto).str.strip().str.lower()
    vazios = texto.isin(_VALORES_AUSENTES)
    saida.loc[vazios, "BASE"] = base
    return saida


def _unificar_nomes_colunas(frames: list[pd.DataFrame]) -> list[pd.DataFrame]:
    """Uma grafia por coluna na junção. O primeiro arquivo define o nome."""
    canonico: dict[str, str] = {}
    unificados: list[pd.DataFrame] = []
    for df in frames:
        renomear: dict[str, str] = {}
        for col in df.columns:
            chave = str(col).casefold()
            nome = canonico.get(chave)
            if nome is None:
                canonico[chave] = str(col)
            elif nome != col:
                renomear[str(col)] = nome
        unificados.append(df.rename(columns=renomear) if renomear else df)
    return unificados


# ── Descoberta e Metadados de Arquivos ──────────────────────────────
def obter_pasta_robo_padrao() -> str:
    base = Path(__file__).resolve().parent
    for cand in (base / "dados", Path.home() / "robo", Path.home() / "Downloads"):
        if cand.is_dir():
            return str(cand.resolve())
    return str(Path.home() / "Downloads")


def formatar_tamanho(bytes_: int) -> str:
    if bytes_ < 1024:
        return f"{bytes_} B"
    if bytes_ < 1024 * 1024:
        return f"{bytes_ / 1024:.1f} KB"
    return f"{bytes_ / (1024 * 1024):.2f} MB"


def obter_metadados_arquivo(caminho: str | None) -> dict[str, str]:
    vazio = {
        "nome": Path(caminho).name if caminho else "Desconhecido",
        "ext": Path(caminho).suffix.lower() if caminho else "",
        "tamanho": "—",
        "modificado": "—",
        "caminho_completo": caminho or "",
    }
    if not caminho or not os.path.exists(caminho):
        return vazio
    try:
        p = Path(caminho)
        stat = p.stat()
        return {
            "nome": p.name,
            "ext": p.suffix.lower(),
            "tamanho": formatar_tamanho(stat.st_size),
            "modificado": datetime.fromtimestamp(stat.st_mtime).strftime(
                "%d/%m/%Y %H:%M:%S"
            ),
            "caminho_completo": str(p.resolve()),
        }
    except OSError:
        return vazio


def _stat_basico(caminho: str) -> tuple[float | None, int | None]:
    try:
        st_ = os.stat(caminho)
        return float(st_.st_mtime), int(st_.st_size)
    except OSError:
        return None, None


def _arquivo_ignorado(nome: str) -> bool:
    baixo = nome.lower()
    if baixo.startswith(("~$", ".")):
        return True
    return baixo.endswith(EXTENSOES_TEMP)


def _listar_arquivos(pasta: Path, recursivo: bool) -> list[Path]:
    try:
        if recursivo:
            return [p for p in pasta.rglob("*") if p.is_file()]
        return [p for p in pasta.iterdir() if p.is_file()]
    except OSError as e:
        logger.warning("Não foi possível listar %s: %s", pasta, e)
        return []


def _casa_padrao(nome: str, padroes: Sequence[str]) -> bool:
    baixo = nome.lower()
    return any(fnmatch.fnmatch(baixo, p.lower()) for p in padroes)


def _listar_candidatos_com_stat(
    pasta: str,
    padroes: tuple[str, ...] = PADROES_TOTALE,
    recursivo: bool = False,
) -> list[Candidato]:
    p = Path(pasta)
    if not p.is_dir():
        return []

    arquivos = [
        arq for arq in _listar_arquivos(p, recursivo) if not _arquivo_ignorado(arq.name)
    ]
    encontrados = [arq for arq in arquivos if _casa_padrao(arq.name, padroes)]

    if not encontrados:
        encontrados = [
            arq
            for arq in arquivos
            if "atividades" in arq.name.lower()
            and arq.suffix.lower() in EXTENSOES_VALIDAS
        ]
    if not encontrados:
        marcadores = tuple(m.lower() for m in MAPA_BASES_ARQUIVO)
        encontrados = [
            arq
            for arq in arquivos
            if arq.suffix.lower() in EXTENSOES_VALIDAS
            and any(m in arq.name.lower() for m in marcadores)
        ]

    candidatos: list[Candidato] = []
    for arq in encontrados:
        mtime, size = _stat_basico(str(arq))
        if mtime is None or size is None or size <= 0:
            continue
        candidatos.append((str(arq), mtime, size))
    candidatos.sort(key=lambda item: item[1], reverse=True)
    return candidatos


def buscar_arquivo_mais_recente(
    pasta: str,
    padroes: tuple[str, ...] = PADROES_TOTALE,
) -> tuple[str | None, float | None, int | None]:
    candidatos = _listar_candidatos_com_stat(pasta, padroes)
    if not candidatos:
        return None, None, None
    caminho, mtime, size = candidatos[0]
    return caminho, mtime, size


def mapear_arquivos_por_base(
    pasta: str, recursivo: bool = False
) -> dict[str, Candidato]:
    por_base: dict[str, Candidato] = {}
    for item in _listar_candidatos_com_stat(pasta, recursivo=recursivo):
        base = detectar_base_arquivo(item[0])
        if base and base not in por_base:
            por_base[base] = item
    return por_base


def selecionar_arquivos(
    candidatos: Sequence[Candidato],
    modo: ModoLeitura = "ultimo",
) -> list[Candidato]:
    """Seleciona somente o arquivo de maior mtime, sem separar por base."""
    del modo
    if not candidatos:
        return []
    return [max(candidatos, key=lambda item: item[1])]


def contar_arquivos_por_base(candidatos: Sequence[Candidato]) -> dict[str, int]:
    contagem: dict[str, int] = {}
    for caminho, _, _ in candidatos:
        base = detectar_base_arquivo(caminho) or "SEM MARCADOR"
        contagem[base] = contagem.get(base, 0) + 1
    return contagem


# ── Leitura Resiliente (CSV / Excel) ────────────────────────────────
def _ler_bytes_estavel(caminho: str, tamanho_maximo: int) -> bytes:
    try:
        tamanho = os.path.getsize(caminho)
    except OSError as e:
        raise ArquivoBloqueadoError(
            f"'{Path(caminho).name}' não pôde ser medido."
        ) from e
    if tamanho > tamanho_maximo:
        raise ValueError(
            f"'{Path(caminho).name}' tem {formatar_tamanho(tamanho)}, "
            f"acima do limite de {formatar_tamanho(tamanho_maximo)}."
        )
    if tamanho == 0:
        raise ValueError(f"'{Path(caminho).name}' está vazio.")
    try:
        with open(caminho, "rb") as fh:
            raw = fh.read()
        tamanho_depois = os.path.getsize(caminho)
    except PermissionError as e:
        raise ArquivoBloqueadoError(
            f"'{Path(caminho).name}' está bloqueado (OneDrive/Excel)."
        ) from e
    except OSError as e:
        raise ArquivoBloqueadoError(f"'{Path(caminho).name}' ficou inacessível.") from e
    if len(raw) != tamanho or tamanho_depois != tamanho:
        raise ArquivoBloqueadoError(
            f"'{Path(caminho).name}' mudou durante a leitura. Aguardando o próximo ciclo."
        )
    return raw


def _validar_xlsx(raw: bytes, nome: str) -> None:
    if len(raw) < _TAMANHO_MINIMO_XLSX:
        raise ValueError(f"'{nome}' tem {len(raw)} bytes — download incompleto.")
    bio = BytesIO(raw)
    if not zipfile.is_zipfile(bio):
        raise ValueError(
            f"'{nome}' não é um .xlsx válido (zip corrompido ou extensão errada)."
        )
    bio.seek(0)
    with zipfile.ZipFile(bio) as zf:
        ruim = zf.testzip()
    if ruim is not None:
        raise ValueError(f"'{nome}' está corrompido no membro interno: {ruim}")


def _ler_excel_via_openpyxl_nativo(
    raw: bytes, sheet_name: str | int, header: int
) -> pd.DataFrame:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        ws = None
        if isinstance(sheet_name, int):
            if 0 <= sheet_name < len(wb.worksheets):
                ws = wb.worksheets[sheet_name]
        elif sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        if ws is None:
            ws = wb.active
        if ws is None and wb.worksheets:
            ws = wb.worksheets[0]
        if ws is None:
            return pd.DataFrame()

        linhas = list(ws.iter_rows(values_only=True))
        if header >= len(linhas):
            return pd.DataFrame()
        header_raw = linhas[header]
        if header_raw is None:
            return pd.DataFrame()

        colunas = _nomes_colunas_unicos(header_raw)
        n_cols = len(colunas)
        data: list[list[object]] = []
        for row in linhas[header + 1 :]:
            valores: list[object] = list(row)
            if len(valores) < n_cols:
                valores.extend([None] * (n_cols - len(valores)))
            elif len(valores) > n_cols:
                valores = valores[:n_cols]
            data.append(valores)
        return pd.DataFrame(data, columns=pd.Index(colunas))
    finally:
        wb.close()


def _ler_excel_robusto(
    caminho: str,
    sheet_name: str | int | None = 0,
    header: int = 0,
    tamanho_maximo: int = _TAMANHO_MAXIMO_PADRAO,
) -> tuple[pd.DataFrame, str]:
    planilha: str | int = 0 if sheet_name is None else sheet_name
    ext = Path(caminho).suffix.lower()
    nome = Path(caminho).name
    raw = _ler_bytes_estavel(caminho, tamanho_maximo)
    if ext == ".xlsx":
        _validar_xlsx(raw, nome)

    erros: list[str] = []

    def _tentar(
        rotulo: str,
        fn: Callable[[], pd.DataFrame | dict[str, pd.DataFrame]],
    ) -> pd.DataFrame | None:
        try:
            bruto = fn()
            if isinstance(bruto, dict):
                df = next(iter(bruto.values())) if len(bruto) else pd.DataFrame()
            else:
                df = bruto
            if df.empty:
                erros.append(f"{rotulo}: planilha vazia")
                return None
            return df
        except PandasOptionError as e:
            erros.append(f"{rotulo}: OptionError {e}")
            return None
        except ImportError as e:
            erros.append(f"{rotulo}: engine ausente ({e})")
            return None
        except ArquivoBloqueadoError:
            raise
        except PermissionError as e:
            raise ArquivoBloqueadoError(str(e)) from e
        except Exception as e:
            erros.append(f"{rotulo}: {type(e).__name__}: {e}")
            logger.debug("Falha ao ler %s via %s", nome, rotulo, exc_info=True)
            return None

    def _via_temp(engine: str, sufixo: str) -> pd.DataFrame:
        fd, tmp = tempfile.mkstemp(suffix=sufixo)
        os.close(fd)
        try:
            with open(tmp, "wb") as out:
                out.write(raw)
            return pd.read_excel(
                tmp,
                sheet_name=planilha,
                header=header,
                dtype=str,
                engine=cast(EngineExcel, engine),
            )
        finally:
            try:
                os.unlink(tmp)
            except OSError:
                pass

    def _via_buffer(engine: str, nome_buffer: str) -> pd.DataFrame:
        bio = _BytesNomeados(raw, nome_buffer)
        return pd.read_excel(
            bio,
            sheet_name=planilha,
            header=header,
            dtype=str,
            engine=cast(EngineExcel, engine),
        )

    tentativas: list[tuple[str, Callable[[], pd.DataFrame]]] = []
    if ext == ".xls":
        tentativas.append(("calamine/.xls", lambda: _via_temp("calamine", ".xls")))
        tentativas.append(("xlrd/.xls", lambda: _via_temp("xlrd", ".xls")))
    else:
        tentativas.extend(
            [
                ("openpyxl/temp.xlsx", lambda: _via_temp("openpyxl", ".xlsx")),
                ("openpyxl/BytesIO", lambda: _via_buffer("openpyxl", "arquivo.xlsx")),
                ("calamine", lambda: _via_buffer("calamine", "arquivo.xlsx")),
                (
                    "openpyxl nativo",
                    lambda: _ler_excel_via_openpyxl_nativo(raw, planilha, header),
                ),
            ]
        )

    for rotulo, fn in tentativas:
        df = _tentar(rotulo, fn)
        if df is not None:
            return df, rotulo

    dica = ""
    joined = " | ".join(erros)
    if "engine ausente" in joined.lower():
        dica = (
            " → pip install openpyxl"
            if ext != ".xls"
            else " → pip install python-calamine"
        )
    raise RuntimeError(
        f"{joined}{dica}" if erros else f"Falha desconhecida ao ler {nome}."
    )


_ENCODINGS_CSV: tuple[str, ...] = ("utf-8-sig", "cp1252", "latin1", "iso-8859-1", "utf-8")
_SEPARADORES_CSV: tuple[str, ...] = (";", ",")


def _ordem_tentativas_csv(raw: bytes) -> list[tuple[str, str]]:
    """Ordena ``(encoding, separador)`` pela probabilidade real do arquivo.

    A leitura ingênua tenta 5 encodings × 2 separadores: um CSV separado por
    vírgula exigia 2 parses completos do arquivo (0,28 s cada em bases grandes).
    Aqui o separador é contado na primeira linha decodificável, então a
    tentativa certa costuma ser a primeira — o restante fica como fallback.
    """
    primeira_linha = ""
    enc_inicial = _ENCODINGS_CSV[0]
    for enc in _ENCODINGS_CSV:
        try:
            primeira_linha = raw[:8192].decode(enc, "strict").split("\n", 1)[0]
            enc_inicial = enc
            break
        except UnicodeDecodeError:
            continue
    if not primeira_linha:
        primeira_linha = raw[:8192].decode("latin1", "ignore").split("\n", 1)[0]

    contagens = {sep: primeira_linha.count(sep) for sep in _SEPARADORES_CSV}
    separadores = sorted(contagens, key=lambda sep: contagens[sep], reverse=True)

    ordem: list[tuple[str, str]] = [(enc_inicial, separadores[0])]
    for sep in separadores:
        if (enc_inicial, sep) not in ordem:
            ordem.append((enc_inicial, sep))
    for enc in _ENCODINGS_CSV:
        for sep in separadores:
            if (enc, sep) not in ordem:
                ordem.append((enc, sep))
    return ordem


def _ler_csv_de_bytes(raw: bytes, nome: str) -> tuple[pd.DataFrame | None, str | None]:
    ultimo: Exception | None = None
    for enc, sep in _ordem_tentativas_csv(raw):
        try:
            df = pd.read_csv(
                BytesIO(raw),
                sep=sep,
                encoding=enc,
                dtype=str,
                low_memory=False,
                on_bad_lines="skip",
            )
        except Exception as e:
            ultimo = e
            continue
        if df.shape[1] <= 1:
            continue
        return df, f"csv {enc} sep={sep!r}"
    detalhe = f" ({type(ultimo).__name__}: {ultimo})" if ultimo else ""
    return None, f"{nome}: falha ao decodificar CSV{detalhe}"


def ler_arquivo_detalhado(
    caminho_arquivo: str,
    colunas_esperadas: list[str] | None = None,
    sheet_name: str | int | None = 0,
    tamanho_maximo: int = _TAMANHO_MAXIMO_PADRAO,
) -> LeituraArquivo:
    if not caminho_arquivo or not os.path.exists(caminho_arquivo):
        nome = Path(caminho_arquivo).name if caminho_arquivo else "—"
        return LeituraArquivo(None, f"Arquivo não encontrado: {nome}", None)

    nome = Path(caminho_arquivo).name
    ext = Path(caminho_arquivo).suffix.lower()
    if ext not in EXTENSOES_VALIDAS:
        return LeituraArquivo(
            None, f"{nome}: extensão {ext or '—'} não é CSV/Excel.", None
        )

    try:
        if ext in (".xlsx", ".xls"):
            df, motor = _ler_excel_robusto(
                caminho_arquivo,
                sheet_name=sheet_name,
                header=0,
                tamanho_maximo=tamanho_maximo,
            )
            df = _padronizar_dataframe(df)
            if _cabecalho_fraco(list(df.columns)) or not _contem_colunas(
                df, colunas_esperadas
            ):
                try:
                    alternativo, motor_alt = _ler_excel_robusto(
                        caminho_arquivo,
                        sheet_name=sheet_name,
                        header=1,
                        tamanho_maximo=tamanho_maximo,
                    )
                except Exception as e:
                    logger.debug("Header alternativo falhou em %s: %s", nome, e)
                else:
                    alternativo = _padronizar_dataframe(alternativo)
                    alt_melhor = (
                        _contem_colunas(alternativo, colunas_esperadas)
                        and not _contem_colunas(df, colunas_esperadas)
                    ) or (
                        _cabecalho_fraco(list(df.columns))
                        and not _cabecalho_fraco(list(alternativo.columns))
                    )
                    if alt_melhor:
                        df, motor = alternativo, f"{motor_alt} header=1"
        else:
            raw = _ler_bytes_estavel(caminho_arquivo, tamanho_maximo)
            df_csv, motor_csv = _ler_csv_de_bytes(raw, nome)
            if df_csv is None:
                return LeituraArquivo(None, motor_csv, None)
            df, motor = _padronizar_dataframe(df_csv), motor_csv
    except ArquivoBloqueadoError:
        raise
    except Exception as e:
        logger.debug("Falha ao ler %s: %s", nome, e, exc_info=True)
        return LeituraArquivo(None, f"{nome}: {type(e).__name__} — {e}", None)

    if df.empty:
        return LeituraArquivo(
            None, f"{nome}: arquivo lido, porém sem linhas de dados.", motor
        )

    df = _alinhar_colunas_esperadas(df, colunas_esperadas)
    faltando = _faltam_colunas(df, colunas_esperadas)
    if faltando:
        return LeituraArquivo(
            None,
            f"{nome}: colunas ausentes → {', '.join(faltando)}",
            motor,
        )
    return LeituraArquivo(df, None, motor)


def processar_candidatos(
    candidatos: Sequence[Candidato],
    modo: ModoLeitura = "ultimo",
    colunas_esperadas: list[str] | None = None,
    sheet_name: str | int | None = 0,
    tamanho_maximo: int = _TAMANHO_MAXIMO_PADRAO,
) -> ResultadoCiclo:
    selecionados = selecionar_arquivos(candidatos, modo)
    if not selecionados:
        return ResultadoCiclo(None, (), (), (), None, (), 0)

    frames: list[pd.DataFrame] = []
    bases: list[str] = []
    arquivos: list[str] = []
    erros: list[str] = []
    motores: list[str] = []
    bloqueios: list[str] = []

    for caminho, _, _ in selecionados:
        try:
            leitura = ler_arquivo_detalhado(
                caminho,
                colunas_esperadas=colunas_esperadas,
                sheet_name=sheet_name,
                tamanho_maximo=tamanho_maximo,
            )
        except ArquivoBloqueadoError as e:
            bloqueios.append(str(e))
            continue
        if leitura.erro or leitura.df is None:
            erros.append(leitura.erro or f"{Path(caminho).name}: leitura vazia")
            continue
        df_parcial = leitura.df
        base = detectar_base_arquivo(caminho)
        if base:
            df_parcial = _aplicar_base_no_dataframe(df_parcial, base)
            if base not in bases:
                bases.append(base)
        frames.append(df_parcial)
        arquivos.append(caminho)
        if leitura.motor:
            motores.append(f"{Path(caminho).name}: {leitura.motor}")

    if not frames:
        return ResultadoCiclo(
            None,
            tuple(bases),
            (),
            tuple(erros),
            bloqueios[0] if bloqueios else None,
            tuple(motores),
            0,
        )

    df = frames[0]
    return ResultadoCiclo(
        df,
        tuple(bases),
        tuple(arquivos),
        tuple(erros),
        bloqueios[0] if bloqueios else None,
        tuple(motores),
        0,
    )


# ── Fachada OO (Compatibilidade com TotaleRoboEngine) ───────────────
class TotaleRoboEngine:
    """Fachada unificada de detecção e leitura de exportações TOTALE."""

    PADROES = PADROES_TOTALE
    MAPA_BASES = MAPA_BASES_ARQUIVO

    _ultimo_erro: str | None = None
    _ultimo_motor: str | None = None

    @classmethod
    def ultimo_erro(cls) -> str | None:
        return cls._ultimo_erro

    @classmethod
    def ultimo_motor(cls) -> str | None:
        return cls._ultimo_motor

    @classmethod
    def _registrar(cls, erro: str | None, motor: str | None = None) -> None:
        cls._ultimo_erro = erro
        cls._ultimo_motor = motor
        if erro:
            logger.debug("%s", erro)

    @staticmethod
    def detectar_base(caminho: str | Path | None) -> str | None:
        return detectar_base_arquivo(caminho)

    @staticmethod
    def listar_validos(pasta: str | Path, *, recursivo: bool = False) -> list[Path]:
        candidatos = _listar_candidatos_com_stat(str(pasta), recursivo=recursivo)
        return [Path(c[0]) for c in candidatos]

    @staticmethod
    def buscar_mais_recente(pasta: str | Path) -> Path | None:
        caminho, _, _ = buscar_arquivo_mais_recente(str(pasta))
        return Path(caminho) if caminho else None

    @staticmethod
    def ler_mais_recente(
        pasta: str | Path,
        *,
        aplicar_base: bool = False,
        sheet_name: str | int = 0,
    ) -> pd.DataFrame | None:
        caminho = TotaleRoboEngine.buscar_mais_recente(pasta)
        if caminho is None:
            TotaleRoboEngine._registrar(f"Nenhum arquivo TOTALE válido em {pasta}")
            return None
        return TotaleRoboEngine.ler_dados(
            caminho,
            aplicar_base=aplicar_base,
            sheet_name=sheet_name,
        )

    @staticmethod
    def arquivos_por_base(pasta: str | Path) -> dict[str, Path]:
        caminho = TotaleRoboEngine.buscar_mais_recente(pasta)
        if caminho is None:
            return {}
        chave = TotaleRoboEngine.detectar_base(caminho) or "SEM MARCADOR"
        return {chave: caminho}

    @staticmethod
    def assinatura(caminho: str | Path) -> tuple[int, int] | None:
        try:
            stat = Path(caminho).stat()
        except OSError:
            return None
        return int(stat.st_size), int(stat.st_mtime_ns)

    @staticmethod
    def verificar_estabilidade(
        caminho: str | Path,
        ciclos: int = 2,
        intervalo: float = 0.8,
    ) -> bool:
        arquivo = Path(caminho)
        if not arquivo.is_file():
            return False
        try:
            anterior = arquivo.stat()
        except OSError:
            return False
        if anterior.st_size <= 0:
            return False
        for _ in range(max(ciclos, 1)):
            time.sleep(intervalo)
            try:
                atual = arquivo.stat()
            except OSError:
                return False
            if (
                atual.st_size != anterior.st_size
                or atual.st_mtime_ns != anterior.st_mtime_ns
            ):
                return False
            anterior = atual
        return anterior.st_size > 0

    @staticmethod
    def ler_dados(
        caminho: str | Path,
        *,
        aplicar_base: bool = False,
        sheet_name: str | int = 0,
    ) -> pd.DataFrame | None:
        try:
            leitura = ler_arquivo_detalhado(
                str(caminho),
                sheet_name=sheet_name,
            )
        except ArquivoBloqueadoError as e:
            TotaleRoboEngine._registrar(str(e))
            return None

        if leitura.erro or leitura.df is None:
            TotaleRoboEngine._registrar(leitura.erro, leitura.motor)
            return None

        df = leitura.df
        if aplicar_base:
            base = TotaleRoboEngine.detectar_base(caminho)
            if base:
                df = _aplicar_base_no_dataframe(df, base)
        TotaleRoboEngine._registrar(None, leitura.motor)
        return df


__all__ = [
    "EXTENSOES_TEMP",
    "EXTENSOES_VALIDAS",
    "MAPA_BASES_ARQUIVO",
    "PADROES_TOTALE",
    "ArquivoBloqueadoError",
    "Candidato",
    "EngineExcel",
    "LeituraArquivo",
    "ModoLeitura",
    "ResultadoCiclo",
    "TotaleRoboEngine",
    "_TAMANHO_MAXIMO_PADRAO",
    "_TAMANHO_MINIMO_XLSX",
    "_VALORES_AUSENTES",
    "_aplicar_base_no_dataframe",
    "_cabecalho_fraco",
    "_celula_para_texto",
    "_listar_candidatos_com_stat",
    "_nomes_colunas_unicos",
    "_ordem_tentativas_csv",
    "_padronizar_coluna_texto",
    "_padronizar_dataframe",
    "_unificar_nomes_colunas",
    "buscar_arquivo_mais_recente",
    "contar_arquivos_por_base",
    "detectar_base_arquivo",
    "detectar_marcador_arquivo",
    "formatar_tamanho",
    "ler_arquivo_detalhado",
    "mapear_arquivos_por_base",
    "obter_metadados_arquivo",
    "obter_pasta_robo_padrao",
    "processar_candidatos",
    "selecionar_arquivos",
]
