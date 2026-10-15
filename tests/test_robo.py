"""Regressões do robô único: arquivos, leitura, CLI e sincronismo Streamlit."""

from __future__ import annotations

import os
import subprocess
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest
from xlwt import Workbook as XlsWorkbook

import robo


@pytest.fixture
def pasta(tmp_path: Path) -> Path:
    return tmp_path


def _csv(pasta: Path, nome: str, texto: str, encoding: str = "utf-8") -> Path:
    arquivo = pasta / nome
    arquivo.write_text(texto, encoding=encoding)
    return arquivo


def _verificar_sem_fragmento(**kwargs: object) -> None:
    # O fragmento é decorado pelo Streamlit só no portal; o teste roda o ciclo
    # original com estado e notificações controlados, sem sessão de navegador.
    assert hasattr(robo._executar_verificacao_robo, "__wrapped__")
    robo._executar_verificacao_robo.__wrapped__(**kwargs)


def test_escolhe_o_mais_novo_da_pasta_inteira_e_mapeia_bases(pasta: Path) -> None:
    antigo = _csv(pasta, "Atividades-SPO.csv", "OS;NOME\n1;Antigo\n")
    novo = _csv(pasta, "GRS.csv", "OS;NOME\n2;Novo\n")
    _csv(pasta, "~$Atividades-SPO.csv", "OS;NOME\n3;Em edição\n")
    _csv(pasta, "Outro.csv", "OS;NOME\n4;Sem marcador\n")
    os.utime(antigo, ns=(1_700_000_000_000_000_000,) * 2)
    os.utime(novo, ns=(1_700_000_001_000_000_000,) * 2)

    candidatos = robo._listar_candidatos_com_stat(str(pasta))
    assert [c[0] for c in candidatos] == [str(novo), str(antigo)]
    assert robo.selecionar_arquivos(candidatos, "todos") == [candidatos[0]]
    assert set(robo.mapear_arquivos_por_base(str(pasta))) == {"NET-LESTE", "NET-GUARULHOS"}
    assert robo.buscar_arquivo_mais_recente(str(pasta)) == (
        str(novo),
        1_700_000_001.0,
        novo.stat().st_size,
    )
    ciclo = robo.executar_ciclo_headless(str(pasta))
    assert ciclo.arquivos == (str(novo),)
    assert ciclo.bases == ("NET-GUARULHOS",)
    assert ciclo.df is not None
    assert ciclo.df.loc[0, "BASE"] == "NET-GUARULHOS"
    assert ciclo.df.loc[0, "OS"] == "2"


def test_subpastas_so_entram_quando_recursivo_e_assinatura_e_precisa(pasta: Path) -> None:
    arquivo = _csv(pasta, "Atividades-SPO.csv", "OS\n1\n")
    momento = 1_700_000_000_000_000_000
    os.utime(arquivo, ns=(momento, momento))
    anterior = robo._listar_candidatos_com_stat(str(pasta))
    arquivo.write_text("OS\n2\n", encoding="utf-8")  # mesmo tamanho
    os.utime(arquivo, ns=(momento + 100, momento + 100))
    atual = robo._listar_candidatos_com_stat(str(pasta))
    assert anterior[0][2] == atual[0][2]
    assert anterior[0][1] != atual[0][1]

    subpasta = pasta / "subpasta"
    subpasta.mkdir()
    _csv(subpasta, "Atividades-ABCDM.csv", "OS\n3\n")
    assert len(robo._listar_candidatos_com_stat(str(pasta))) == 1
    assert len(robo._listar_candidatos_com_stat(str(pasta), recursivo=True)) == 2


@pytest.mark.parametrize(
    ("nome", "conteudo", "colunas", "valor"),
    [
        ("Atividades-SPO.csv", b"OS;NOME\n123;Maria\n", ["OS", "NOME"], "Maria"),
        ("Atividades-SPO.csv", "OS;NOME\n456;Jos\xe9\n".encode("cp1252"), ["OS", "NOME"], "José"),
        (
            "Atividades-SPO.csv",
            b'"Descricao;ativa",OS\n"Chamado;ok",789\n',
            ["Descricao;ativa", "OS"],
            "789",
        ),
        ("Atividades-SPO.csv", b"OS\n101\n", ["OS"], "101"),
        ("Atividades-SPO.csv", b"OS\tNOME\n102\tJulia\n", ["OS", "NOME"], "Julia"),
    ],
)
def test_csv_leitura_resiliente(
    pasta: Path, nome: str, conteudo: bytes, colunas: list[str], valor: str
) -> None:
    caminho = pasta / nome
    caminho.write_bytes(conteudo)
    leitura = robo.ler_arquivo_detalhado(str(caminho), colunas_esperadas=colunas)
    assert leitura.erro is None
    assert leitura.df is not None
    assert list(leitura.df.columns) == colunas
    assert valor in leitura.df.iloc[0].tolist()
    assert leitura.motor and leitura.motor.startswith("csv ")


def test_planilha_excel_com_cabecalho_na_segunda_linha_e_base(pasta: Path) -> None:
    caminho = pasta / "Atividades-SPO.xlsx"
    livro = Workbook()
    aba = livro.active
    aba.append(["Relatório de atividades", None])
    aba.append(["OS", "BASE"])
    aba.append([123, None])
    livro.save(caminho)

    leitura = robo.ler_arquivo_detalhado(str(caminho), ["os", "BASE"])
    assert leitura.erro is None
    assert leitura.df is not None
    assert list(leitura.df.columns) == ["os", "BASE"]
    assert leitura.df.loc[0, "os"] == "123"
    assert "header=1" in (leitura.motor or "")
    ciclo = robo.executar_ciclo_headless(str(pasta), colunas_esperadas=["OS", "BASE"])
    assert ciclo.df is not None
    assert ciclo.df.loc[0, "BASE"] == "NET-LESTE"


def test_planilha_xls_com_xlrd_e_xlsm_com_openpyxl(pasta: Path) -> None:
    caminho_xls = pasta / "Atividades-GRS.xls"
    livro_xls = XlsWorkbook()
    aba = livro_xls.add_sheet("Dados")
    for col, valor in enumerate(("OS", "NOME")):
        aba.write(0, col, valor)
    for col, valor in enumerate((99, "Ana")):
        aba.write(1, col, valor)
    livro_xls.save(str(caminho_xls))

    leitura_xls = robo.ler_arquivo_detalhado(str(caminho_xls))
    assert leitura_xls.erro is None
    assert leitura_xls.df is not None
    assert leitura_xls.df.loc[0, "NOME"] == "Ana"
    assert "xlrd" in (leitura_xls.motor or "")

    caminho_xlsm = pasta / "Atividades-ABCDM.xlsm"
    livro_xlsm = Workbook()
    livro_xlsm.active.append(["OS", "NOME"])
    livro_xlsm.active.append([101, "Bia"])
    livro_xlsm.save(caminho_xlsm)
    leitura_xlsm = robo.ler_arquivo_detalhado(str(caminho_xlsm))
    assert leitura_xlsm.erro is None
    assert leitura_xlsm.df is not None
    assert leitura_xlsm.df.loc[0, "NOME"] == "Bia"


def test_aba_inexistente_nao_usa_aba_ativa_silenciosamente(pasta: Path) -> None:
    arquivo = pasta / "Atividades-SPO.xlsx"
    livro = Workbook()
    livro.active.append(["OS", "NOME"])
    livro.active.append([1, "Ana"])
    livro.save(arquivo)

    with pytest.raises(ValueError, match="Aba 'Inexistente' não encontrada"):
        robo._ler_excel_via_openpyxl_nativo(arquivo.read_bytes(), "Inexistente", 0)
    leitura = robo.ler_arquivo_detalhado(str(arquivo), sheet_name="Inexistente")
    assert leitura.df is None
    assert "Inexistente" in (leitura.erro or "")


def test_nomes_unicos_e_numeros_grandes_nao_estouram_int64() -> None:
    assert robo._nomes_colunas_unicos(["A", "A", "A.1", " ", "A"]) == [
        "A",
        "A.2",
        "A.1",
        "col_3",
        "A.3",
    ]
    serie = pd.Series([1.0, float(2**64), None])
    padronizada = robo._padronizar_coluna_texto(serie)
    assert padronizada.iloc[0] == "1"
    assert padronizada.iloc[1] == str(2**64)
    assert pd.isna(padronizada.iloc[2])


def test_arquivo_alterado_durante_leitura_e_tentado_novamente(
    pasta: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    arquivo = _csv(pasta, "Atividades-SPO.csv", "OS\n1\n")
    stat_real = os.stat
    chamadas = 0

    def stat_alterado(caminho: str, *args: object, **kwargs: object) -> object:
        nonlocal chamadas
        info = stat_real(caminho, *args, **kwargs)
        if str(caminho) == str(arquivo):
            chamadas += 1
            if chamadas == 2:
                return SimpleNamespace(
                    st_size=info.st_size,
                    st_mtime_ns=info.st_mtime_ns + 1,
                    st_dev=info.st_dev,
                    st_ino=info.st_ino,
                )
        return info

    monkeypatch.setattr(robo.os, "stat", stat_alterado)
    with pytest.raises(robo.ArquivoBloqueadoError, match="mudou durante a leitura"):
        robo._ler_bytes_estavel(str(arquivo), 1000)


def test_csv_malformado_nao_perde_coluna_ao_inferir_indice(pasta: Path) -> None:
    arquivo = _csv(pasta, "Atividades-SPO.csv", "OS;NOME\n1;Ana;extra\n")
    leitura = robo.ler_arquivo_detalhado(str(arquivo))
    assert leitura.df is None
    assert "ParserWarning" in (leitura.erro or "")


def test_erros_de_leitura_sao_explicitos(pasta: Path) -> None:
    inexistente = robo.ler_arquivo_detalhado(str(pasta / "inexistente.csv"))
    assert inexistente.df is None and "não encontrado" in (inexistente.erro or "")
    txt = _csv(pasta, "Atividades-SPO.txt", "OS\n1\n")
    assert "extensão" in (robo.ler_arquivo_detalhado(str(txt)).erro or "")
    grande = _csv(pasta, "Atividades-SPO.csv", "OS\n123456\n")
    assert "acima do limite" in (
        robo.ler_arquivo_detalhado(str(grande), tamanho_maximo=5).erro or ""
    )
    corrompido = pasta / "Atividades-SPO.xlsx"
    corrompido.write_bytes(b"not an Excel file" * 300)
    assert "zip" in (robo.ler_arquivo_detalhado(str(corrompido)).erro or "").lower()


def test_cli_headless_grava_utf8_sem_sobrescrever_a_origem(pasta: Path) -> None:
    entrada = _csv(pasta, "Atividades-GRS.csv", "OS;NOME\n1;José\n")
    destino = pasta / "saida.csv"
    comando = subprocess.run(
        [sys.executable, str(Path(robo.__file__)), str(pasta), str(destino)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert comando.returncode == 0, comando.stderr
    assert "1 linhas" in comando.stdout
    assert destino.read_bytes().startswith(b"\xef\xbb\xbf")
    assert pd.read_csv(destino, sep=";").loc[0, "NOME"] == "José"
    conteudo_original = entrada.read_bytes()
    with pytest.raises(ValueError, match="não pode sobrescrever"):
        robo.executar_ciclo_headless(str(pasta), destino_csv=str(entrada))
    tentativa = subprocess.run(
        [sys.executable, str(Path(robo.__file__)), str(pasta), str(entrada)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert tentativa.returncode == 2
    assert "não pode sobrescrever" in tentativa.stderr
    assert "Traceback" not in tentativa.stderr
    assert entrada.read_bytes() == conteudo_original


def test_etl_falhou_nao_marca_processado_e_reexecuta(
    pasta: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _csv(pasta, "Atividades-SPO.csv", "OS;NOME\n1;Ana\n")
    anterior = pd.DataFrame({"OS": ["anterior"]})
    estado: dict[str, object] = {
        "df_memoria_raw": anterior,
        "df_memoria": anterior,
        "origem_dados": "manual",
    }
    monkeypatch.setattr(
        robo,
        "st",
        SimpleNamespace(session_state=estado, toast=lambda *a, **kw: None, rerun=lambda **kw: None),
    )

    def etl_com_falha(df: pd.DataFrame, _gs: pd.DataFrame) -> pd.DataFrame:
        df.loc[0, "OS"] = "alterado"
        raise ValueError("ETL falhou")

    _verificar_sem_fragmento(
        pasta_monitorada=str(pasta), etl_fn=etl_com_falha, ciclos_estabilidade=1
    )
    assert "ETL falhou" in str(estado["robo_erro"])
    assert estado["df_memoria"] is anterior
    assert estado["df_memoria_raw"] is anterior
    assert estado["origem_dados"] == "manual"
    assert "robo_processado_sig" not in estado
    assert estado["robo_processando"] is False

    def etl_ok(df: pd.DataFrame, _gs: pd.DataFrame) -> pd.DataFrame:
        df.loc[0, "OS"] = "processado"
        return df

    _verificar_sem_fragmento(pasta_monitorada=str(pasta), etl_fn=etl_ok, ciclos_estabilidade=1)
    assert estado["robo_processado_sig"]
    assert isinstance(estado["df_memoria_raw"], pd.DataFrame)
    assert estado["df_memoria_raw"].loc[0, "OS"] == "1"
    assert estado["df_memoria"].loc[0, "OS"] == "processado"
    assert estado["origem_dados"].startswith("Robô Local")
    assert "robo_erro" not in estado


def test_mudanca_durante_etl_aguarda_novo_ciclo(
    pasta: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    arquivo = _csv(pasta, "Atividades-SPO.csv", "OS\n1\n")
    estado: dict[str, object] = {}
    monkeypatch.setattr(
        robo,
        "st",
        SimpleNamespace(session_state=estado, toast=lambda *a, **kw: None, rerun=lambda **kw: None),
    )

    def etl_altera_fonte(df: pd.DataFrame, _gs: pd.DataFrame) -> pd.DataFrame:
        anterior_ns = arquivo.stat().st_mtime_ns
        arquivo.write_text("OS\n2\n", encoding="utf-8")
        os.utime(arquivo, ns=(anterior_ns + 1_000, anterior_ns + 1_000))
        return df

    _verificar_sem_fragmento(
        pasta_monitorada=str(pasta), etl_fn=etl_altera_fonte, ciclos_estabilidade=1
    )
    assert "robo_processado_sig" not in estado
    assert "df_memoria_raw" not in estado
    assert estado["robo_candidato_stable"] == 0
    _verificar_sem_fragmento(pasta_monitorada=str(pasta), ciclos_estabilidade=1)
    assert estado["df_memoria_raw"].loc[0, "OS"] == "2"


def test_bloqueio_temporario_reprocessa_no_ciclo_seguinte(
    pasta: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _csv(pasta, "Atividades-SPO.csv", "OS\n1\n")
    estado: dict[str, object] = {}
    monkeypatch.setattr(
        robo,
        "st",
        SimpleNamespace(session_state=estado, toast=lambda *a, **kw: None, rerun=lambda **kw: None),
    )
    leitor_real = robo._ler_bytes_estavel
    chamadas = 0

    def leitor_bloqueado(caminho: str, tamanho: int) -> bytes:
        nonlocal chamadas
        chamadas += 1
        if chamadas == 1:
            raise robo.ArquivoBloqueadoError("Arquivo em gravação")
        return leitor_real(caminho, tamanho)

    monkeypatch.setattr(robo, "_ler_bytes_estavel", leitor_bloqueado)
    _verificar_sem_fragmento(pasta_monitorada=str(pasta), ciclos_estabilidade=1)
    assert estado["_robo_aguardando_lock"] == "Arquivo em gravação"
    assert "robo_processado_sig" not in estado
    _verificar_sem_fragmento(pasta_monitorada=str(pasta), ciclos_estabilidade=1)
    assert estado["robo_processado_sig"]
    assert "_robo_aguardando_lock" not in estado


def test_gsheets_nao_e_carregado_sem_etl_e_none_nao_mantem_dado_antigo(
    pasta: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _csv(pasta, "Atividades-SPO.csv", "OS\n2\n")
    estado: dict[str, object] = {"df_memoria": pd.DataFrame({"OS": ["antigo"]})}
    monkeypatch.setattr(
        robo,
        "st",
        SimpleNamespace(session_state=estado, toast=lambda *a, **kw: None, rerun=lambda **kw: None),
    )

    def gsheets() -> pd.DataFrame:
        raise AssertionError("gsheets_fn só deve ser chamada para o ETL")

    _verificar_sem_fragmento(pasta_monitorada=str(pasta), gsheets_fn=gsheets, ciclos_estabilidade=1)
    assert estado["df_memoria"].loc[0, "OS"] == "2"
    estado.pop("robo_processado_sig")
    _verificar_sem_fragmento(
        pasta_monitorada=str(pasta), etl_fn=lambda _df, _gs: None, ciclos_estabilidade=1
    )
    assert estado["df_memoria"].loc[0, "OS"] == "2"


def test_sidebar_streamlit_real_processa_e_atualiza_arquivo(pasta: Path) -> None:
    arquivo = _csv(pasta, "Atividades-SPO.csv", "OS;NOME\n1;Ana\n")
    codigo = (
        "from robo import renderizar_robo_local\n"
        f"renderizar_robo_local(pasta_padrao={str(pasta)!r}, "
        "mostrar_toggle=False, mostrar_config=False, ciclos_estabilidade=1)\n"
    )
    app = AppTest.from_string(codigo, default_timeout=30).run()
    assert not list(app.exception)
    assert app.session_state["df_memoria_raw"].loc[0, "OS"] == "1"
    assert app.session_state["robo_bases_carregadas"] == ["NET-LESTE"]

    antigo_ns = arquivo.stat().st_mtime_ns
    arquivo.write_text("OS;NOME\n2;Ana\n", encoding="utf-8")
    os.utime(arquivo, ns=(antigo_ns + 1_000, antigo_ns + 1_000))
    app.run()
    assert not list(app.exception)
    assert app.session_state["df_memoria_raw"].loc[0, "OS"] == "2"


def test_lock_expirado_ou_invalido_e_liberado(monkeypatch: pytest.MonkeyPatch) -> None:
    estado: dict[str, object] = {"robo_processando": True}
    monkeypatch.setattr(robo, "st", SimpleNamespace(session_state=estado))
    assert not robo._lock_ativo()
    assert estado["robo_processando"] is False
    estado.update(
        robo_processando=True,
        robo_processando_desde=datetime.now(timezone.utc) - timedelta(minutes=4),
    )
    assert not robo._lock_ativo()
    estado.update(
        robo_processando=True,
        robo_processando_desde=datetime.now(timezone.utc),
    )
    assert robo._lock_ativo()
    estado["robo_processando_desde"] = datetime.now(timezone.utc) + timedelta(minutes=4)
    assert not robo._lock_ativo()
